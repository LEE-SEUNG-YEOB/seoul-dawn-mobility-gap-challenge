"""Build bounded spatial inputs for direct-route screening.

This module links demand-candidate endpoint polygons to route-stop records.  It
does not establish walking access, timetable availability, or a service gap.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pyproj import Transformer
from shapely.geometry import Point, shape
from shapely.ops import transform
from shapely.strtree import STRtree


BUFFER_METERS = 500.0
METRIC_CRS = "EPSG:5179"
SOURCE_DATASET_IDS = "OA-22300|OA-1095|data-go-15129688"

OUTPUT_COLUMNS = [
    "candidate_id",
    "endpoint_role",
    "admi_cd",
    "demand_rank",
    "time_bin_start",
    "time_bin_end",
    "route_id",
    "stop_id",
    "ars_id",
    "stop_sequence",
    "route_name",
    "stop_name",
    "longitude",
    "latitude",
    "inside_or_buffer",
    "distance_to_polygon_m",
    "boundary_id",
    "boundary_date",
    "buffer_m",
    "route_present_at_both_endpoints",
    "direction_sequence_compatible",
    "source_dataset_ids",
    "uncertainty_flag",
]

UNKNOWN_COLUMNS = [
    "record_type",
    "candidate_id",
    "endpoint_role",
    "admi_cd",
    "route_id",
    "stop_id",
    "ars_id",
    "boundary_id",
    "boundary_district",
    "boundary_dong_name",
    "reason",
    "source_dataset_ids",
]


def _text_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _float_or_none(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number


def spatial_relation(polygon: Any, point: Point, buffer_m: float = BUFFER_METERS) -> tuple[str | None, float]:
    """Classify a metric point against a metric polygon and screening buffer."""
    distance = float(polygon.distance(point))
    if polygon.covers(point):
        return "inside", 0.0
    if distance <= buffer_m:
        return "buffer_500m", distance
    return None, distance


def coordinate_statuses(rows: Iterable[dict[str, Any]]) -> dict[str, str]:
    """Classify stop IDs as unique, missing, or spatially ambiguous."""
    observed: dict[str, set[tuple[float, float]]] = defaultdict(set)
    missing: set[str] = set()
    for row in rows:
        stop_id = str(row.get("stop_id", ""))
        longitude = _float_or_none(row.get("longitude"))
        latitude = _float_or_none(row.get("latitude"))
        if not stop_id or longitude is None or latitude is None or not (124 <= longitude <= 132 and 33 <= latitude <= 39):
            missing.add(stop_id)
            continue
        observed[stop_id].add((round(longitude, 7), round(latitude, 7)))
    statuses: dict[str, str] = {}
    for stop_id in set(observed) | missing:
        if stop_id in missing and stop_id not in observed:
            statuses[stop_id] = "unknown_missing_coordinate"
        elif len(observed[stop_id]) != 1 or stop_id in missing:
            statuses[stop_id] = "unknown_ambiguous_coordinate"
        else:
            statuses[stop_id] = "unique"
    return statuses


def _read_oa1095(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    source_rows = sheet.iter_rows(values_only=True)
    header = [_text_id(value) for value in next(source_rows)]
    col = {name: index for index, name in enumerate(header)}
    required = {"ROUTE_ID", "노선명", "순번", "NODE_ID", "ARS_ID", "정류소명", "X좌표", "Y좌표"}
    if not required.issubset(col):
        workbook.close()
        raise ValueError(f"OA-1095 missing columns: {sorted(required - set(col))}")
    result: list[dict[str, Any]] = []
    for source_row_number, row in enumerate(source_rows, start=2):
        result.append(
            {
                "source_row_number": source_row_number,
                "route_id": _text_id(row[col["ROUTE_ID"]]),
                "route_name": _text_id(row[col["노선명"]]),
                "stop_sequence": _text_id(row[col["순번"]]),
                "stop_id": _text_id(row[col["NODE_ID"]]),
                "ars_id": _text_id(row[col["ARS_ID"]]),
                "stop_name": _text_id(row[col["정류소명"]]),
                "longitude": _float_or_none(row[col["X좌표"]]),
                "latitude": _float_or_none(row[col["Y좌표"]]),
            }
        )
    workbook.close()
    return result


def _read_candidates(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    required = {"candidate_id", "origin_id", "destination_id", "demand_rank", "time_bin_start", "time_bin_end"}
    if not rows or not required.issubset(rows[0]):
        raise ValueError("Demand candidate schema is incomplete")
    return rows


def _load_boundaries(path: Path) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    document = json.loads(path.read_text(encoding="utf-8"))
    project = Transformer.from_crs("EPSG:4326", METRIC_CRS, always_xy=True).transform
    matched: dict[str, dict[str, Any]] = {}
    unknown: list[dict[str, Any]] = []
    for feature in document["features"]:
        properties = feature["properties"]
        mobility_id = properties.get("mobility_id")
        record = {
            "polygon": transform(project, shape(feature["geometry"])),
            "boundary_id": properties.get("boundary_id"),
            "boundary_date": properties.get("boundary_date"),
            "district": properties.get("district"),
            "dong_name": properties.get("dong_name"),
            "match_method": properties.get("match_method"),
        }
        if mobility_id is None:
            unknown.append(record)
        elif mobility_id in matched:
            raise ValueError(f"Multiple boundary polygons for mobility ID {mobility_id}")
        else:
            matched[str(mobility_id)] = record
    return matched, unknown


def _write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, columns, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def build(root: Path, buffer_m: float = BUFFER_METERS) -> dict[str, Any]:
    root = root.resolve()
    candidate_path = root / "outputs/intermediate/demand/demand_candidates.csv"
    boundary_path = root / "outputs/final/maps/dong_demand.geojson"
    oa1095_path = next((root / "data/raw/bus/OA-1095").glob("*.xlsx"))
    output_dir = root / "outputs/intermediate/demand"
    output_path = output_dir / "direct_service_candidate_stops.csv"
    unknown_path = output_dir / "direct_service_unknowns.csv"
    summary_path = output_dir / "direct_service_summary.json"
    notes_path = output_dir / "direct_service_notes.md"

    candidates = _read_candidates(candidate_path)
    boundaries, unknown_boundaries = _load_boundaries(boundary_path)
    stop_rows = _read_oa1095(oa1095_path)
    statuses = coordinate_statuses(stop_rows)

    unknown_rows: list[dict[str, Any]] = []
    valid_by_stop: dict[str, list[dict[str, Any]]] = defaultdict(list)
    stop_coordinates: dict[str, tuple[float, float]] = {}
    for row in stop_rows:
        status = statuses.get(row["stop_id"], "unknown_missing_coordinate")
        if status != "unique":
            unknown_rows.append(
                {
                    "record_type": "stop_coordinate",
                    "route_id": row["route_id"],
                    "stop_id": row["stop_id"],
                    "ars_id": row["ars_id"],
                    "reason": status,
                    "source_dataset_ids": "OA-1095",
                }
            )
            continue
        valid_by_stop[row["stop_id"]].append(row)
        stop_coordinates[row["stop_id"]] = (float(row["longitude"]), float(row["latitude"]))

    to_metric = Transformer.from_crs("EPSG:4326", METRIC_CRS, always_xy=True)
    stop_ids = sorted(valid_by_stop)
    points = [Point(*to_metric.transform(*stop_coordinates[stop_id])) for stop_id in stop_ids]
    tree = STRtree(points)

    endpoint_matches: dict[str, list[dict[str, Any]]] = {}
    endpoint_ids = sorted({value for row in candidates for value in (row["origin_id"], row["destination_id"])})
    for admi_cd in endpoint_ids:
        boundary = boundaries.get(admi_cd)
        if boundary is None:
            endpoint_matches[admi_cd] = []
            continue
        polygon = boundary["polygon"]
        matches: list[dict[str, Any]] = []
        for point_index in tree.query(polygon.buffer(buffer_m), predicate="intersects"):
            index = int(point_index)
            relation, distance = spatial_relation(polygon, points[index], buffer_m)
            if relation is None:
                continue
            stop_id = stop_ids[index]
            for stop_row in valid_by_stop[stop_id]:
                matches.append(
                    stop_row
                    | {
                        "inside_or_buffer": relation,
                        "distance_to_polygon_m": round(distance, 3),
                        "boundary_id": boundary["boundary_id"],
                        "boundary_date": boundary["boundary_date"],
                    }
                )
        matches.sort(
            key=lambda row: (
                0 if row["inside_or_buffer"] == "inside" else 1,
                row["distance_to_polygon_m"],
                row["route_id"],
                int(row["stop_sequence"]) if row["stop_sequence"].isdigit() else 10**9,
                row["stop_id"],
            )
        )
        endpoint_matches[admi_cd] = matches

    output_rows: list[dict[str, Any]] = []
    candidates_with_common_route = 0
    compatible_candidate_routes = 0
    for candidate in sorted(candidates, key=lambda row: int(row["demand_rank"])):
        origin_matches = endpoint_matches[candidate["origin_id"]]
        destination_matches = endpoint_matches[candidate["destination_id"]]
        origin_by_route: dict[str, list[int]] = defaultdict(list)
        destination_by_route: dict[str, list[int]] = defaultdict(list)
        for row in origin_matches:
            if row["stop_sequence"].isdigit():
                origin_by_route[row["route_id"]].append(int(row["stop_sequence"]))
        for row in destination_matches:
            if row["stop_sequence"].isdigit():
                destination_by_route[row["route_id"]].append(int(row["stop_sequence"]))
        common_routes = set(origin_by_route) & set(destination_by_route)
        compatible_routes = {
            route_id
            for route_id in common_routes
            if any(origin_sequence < destination_sequence for origin_sequence in origin_by_route[route_id] for destination_sequence in destination_by_route[route_id])
        }
        candidates_with_common_route += int(bool(common_routes))
        compatible_candidate_routes += len(compatible_routes)

        for endpoint_role, admi_cd, matches in (
            ("origin", candidate["origin_id"], origin_matches),
            ("destination", candidate["destination_id"], destination_matches),
        ):
            boundary = boundaries.get(admi_cd)
            if boundary is None:
                output_rows.append(
                    {
                        "candidate_id": candidate["candidate_id"],
                        "endpoint_role": endpoint_role,
                        "admi_cd": admi_cd,
                        "demand_rank": candidate["demand_rank"],
                        "time_bin_start": candidate["time_bin_start"],
                        "time_bin_end": candidate["time_bin_end"],
                        "inside_or_buffer": "unknown_boundary",
                        "buffer_m": buffer_m,
                        "source_dataset_ids": SOURCE_DATASET_IDS,
                        "uncertainty_flag": "boundary crosswalk unknown; no stop assignment",
                    }
                )
                unknown_rows.append(
                    {
                        "record_type": "candidate_endpoint",
                        "candidate_id": candidate["candidate_id"],
                        "endpoint_role": endpoint_role,
                        "admi_cd": admi_cd,
                        "reason": "unknown_boundary_crosswalk",
                        "source_dataset_ids": "OA-22300|data-go-15129688",
                    }
                )
                continue
            for match in matches:
                route_id = match["route_id"]
                output_rows.append(
                    {
                        "candidate_id": candidate["candidate_id"],
                        "endpoint_role": endpoint_role,
                        "admi_cd": admi_cd,
                        "demand_rank": candidate["demand_rank"],
                        "time_bin_start": candidate["time_bin_start"],
                        "time_bin_end": candidate["time_bin_end"],
                        **match,
                        "buffer_m": buffer_m,
                        "route_present_at_both_endpoints": str(route_id in common_routes).lower(),
                        "direction_sequence_compatible": str(route_id in compatible_routes).lower() if route_id in common_routes else "",
                        "source_dataset_ids": SOURCE_DATASET_IDS,
                        "uncertainty_flag": "spatial screening only; 2025Q2 boundary and 2026-09-02 route-stop snapshot; no walking path, service calendar, or timetable test",
                    }
                )

    output_rows.sort(
        key=lambda row: (
            int(row["demand_rank"]),
            0 if row["endpoint_role"] == "origin" else 1,
            0 if row["inside_or_buffer"] == "inside" else 1,
            row.get("distance_to_polygon_m", 10**9),
            row.get("route_id", ""),
            row.get("stop_id", ""),
        )
    )
    for boundary in unknown_boundaries:
        unknown_rows.append(
            {
                "record_type": "boundary_feature",
                "admi_cd": "",
                "boundary_id": boundary["boundary_id"],
                "boundary_district": boundary["district"],
                "boundary_dong_name": boundary["dong_name"],
                "reason": f"unknown_name_crosswalk:{boundary['boundary_id']}",
                "source_dataset_ids": "data-go-15129688",
            }
        )
    unknown_rows.sort(key=lambda row: tuple(str(row.get(column, "")) for column in UNKNOWN_COLUMNS))
    _write_csv(output_path, OUTPUT_COLUMNS, output_rows)
    _write_csv(unknown_path, UNKNOWN_COLUMNS, unknown_rows)

    inside_rows = sum(row["inside_or_buffer"] == "inside" for row in output_rows)
    buffer_rows = sum(row["inside_or_buffer"] == "buffer_500m" for row in output_rows)
    candidate_endpoint_pairs = {(row["candidate_id"], row["endpoint_role"]) for row in output_rows if row["inside_or_buffer"] in {"inside", "buffer_500m"}}
    summary = {
        "status": "complete_limited_spatial_screening_input",
        "candidate_count": len(candidates),
        "unique_candidate_endpoint_dongs": len(endpoint_ids),
        "candidate_endpoint_pairs_with_stop_matches": len(candidate_endpoint_pairs),
        "candidate_endpoint_pairs_expected": len(candidates) * 2,
        "output_rows": len(output_rows),
        "inside_rows": inside_rows,
        "buffer_500m_rows": buffer_rows,
        "candidates_with_route_present_at_both_endpoints": candidates_with_common_route,
        "direction_sequence_compatible_candidate_route_pairs": compatible_candidate_routes,
        "oa1095_source_rows": len(stop_rows),
        "oa1095_unique_coordinate_stop_ids": len(valid_by_stop),
        "oa1095_unknown_coordinate_rows": sum(row["record_type"] == "stop_coordinate" for row in unknown_rows),
        "unknown_boundary_features": len(unknown_boundaries),
        "candidate_endpoints_with_unknown_boundary": sum(row["record_type"] == "candidate_endpoint" for row in unknown_rows),
        "buffer_m": buffer_m,
        "metric_crs": METRIC_CRS,
        "source_dataset_ids": SOURCE_DATASET_IDS,
        "output_sha256": _sha256(output_path),
        "unknowns_sha256": _sha256(unknown_path),
        "interpretation": "Administrative-dong spatial screening input only; not an individual walking threshold, direct service finding, or service-gap finding.",
    }
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    notes_path.write_text(
        """# Direct-service spatial screening input

Each row links one top-100 demand-candidate endpoint to one OA-1095 route-stop record whose stop point lies inside the endpoint administrative-dong polygon or no more than 500 metres outside it. Polygons and points are projected from WGS84 to EPSG:5179 before distance and buffering. `inside` has distance 0; `buffer_500m` retains the measured polygon distance.

The 500 m band is an administrative-dong screening buffer. It is not a claim about an individual's walking distance or network access. `route_present_at_both_endpoints` only means the same OA-1095 route ID has spatially selected stops at both endpoint polygons. `direction_sequence_compatible` only checks whether at least one selected origin stop sequence precedes one selected destination stop sequence. Neither field applies service dates, first-service times, actual walking paths, transfers, operations, or observed journeys.

OA-1095 is a 2026-09-02 snapshot, while the validated SGIS boundary is 2025 Q2 and the mobility codebook is August 2026. Boundary matching relies on the audited unique district-and-dong-name crosswalk. Unknown/split boundary features and missing or ambiguous stop coordinates are retained in `direct_service_unknowns.csv`; they are never treated as absent service. This output is a limited spatial screening input under `join_ready=false`, not a service-gap ranking.
""",
        encoding="utf-8",
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--buffer-m", type=float, default=BUFFER_METERS)
    args = parser.parse_args()
    print(json.dumps(build(args.root, args.buffer_m), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

