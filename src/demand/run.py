from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable

import duckdb


PURPOSE_LABELS = {
    "1": "commuting",
    "2": "school",
    "3": "home",
    "4": "shopping",
    "5": "tourism",
    "6": "hospital",
    "7": "other",
}
MODE_LABELS = {
    "1": "air",
    "2": "rail",
    "3": "express_bus",
    "4": "metropolitan_bus",
    "5": "ordinary_bus",
    "6": "subway",
    "7": "walking",
    "8": "vehicle",
    "9": "other",
}
REQUIRED_CANDIDATE_COLUMNS = [
    "candidate_id",
    "origin_id",
    "destination_id",
    "time_bin_start",
    "time_bin_end",
    "demand_value",
    "demand_unit",
    "demand_rank",
    "selection_rule",
    "uncertainty_flag",
    "source_dataset_ids",
]
OPTIONAL_CANDIDATE_COLUMNS = [
    "origin_name",
    "origin_district_name",
    "destination_name",
    "destination_district_name",
]


def normalize_hour(value: str | None) -> int | None:
    """Return a dawn hourly code without confusing HH with HHMM."""
    if value is None:
        return None
    text = str(value).strip()
    if text not in {"3", "03", "4", "04", "5", "05"}:
        return None
    return int(text)


def day_type(yyyymmdd: str) -> str:
    weekday = datetime.strptime(str(yyyymmdd), "%Y%m%d").weekday()
    return "weekday" if weekday < 5 else "weekend"


def _write_csv(path: Path, columns: list[str], rows: Iterable[Iterable[object]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(columns)
        for row in rows:
            writer.writerow(row)
            count += 1
    return count


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _csv_sql_path(path: Path) -> str:
    return path.resolve().as_posix().replace("'", "''")


def _export_query(con: duckdb.DuckDBPyConnection, out: Path, query: str) -> int:
    target = _csv_sql_path(out)
    con.execute(f"COPY ({query}) TO '{target}' (HEADER, DELIMITER ',', FORCE_QUOTE *)")
    return int(con.execute(f"SELECT count(*) FROM ({query}) q").fetchone()[0])


def _build_mobility_aggregates(con: duckdb.DuckDBPyConnection, root: Path) -> None:
    codebook = root / "data/reference/mobility_codebooks/전국_행정동_코드정보/SEOUL_ADMI/ADMI_202608.csv"
    purpose_glob = root / "data/raw/mobility/OA-22300/*.csv"
    mode_glob = root / "data/raw/mobility/OA-22657/*.csv"
    con.execute(
        f"""
        CREATE OR REPLACE TABLE seoul_codes AS
        SELECT DISTINCT trim(ADMI_CD) AS code, ADMI_NM AS name, SGG_NM AS district_name,
                        FULL_NM AS full_name
        FROM read_csv_auto('{_csv_sql_path(codebook)}', header=true, all_varchar=true)
        WHERE SIDO_NM = '서울특별시' AND length(trim(ADMI_CD)) = 8
        """
    )
    common = """
        CASE WHEN dayofweek(strptime(etl_ymd, '%Y%m%d')) BETWEEN 1 AND 5
             THEN 'weekday' ELSE 'weekend' END AS day_type,
        trim(o_admdong_cd) AS origin_id,
        trim(d_admdong_cd) AS destination_id,
        CAST(trim(st_time_cd) AS INTEGER) AS hour,
        sum(try_cast(cnt AS DOUBLE)) AS demand_value,
        count(*) AS source_rows,
        count(try_cast(cnt AS DOUBLE)) AS nonmissing_cnt_rows
    """
    con.execute(
        f"""
        CREATE OR REPLACE TABLE purpose_dawn AS
        SELECT {common}, trim(move_purpose) AS category_code
        FROM read_csv_auto('{_csv_sql_path(purpose_glob)}', header=true,
                           all_varchar=true, union_by_name=true, filename=false)
        WHERE trim(st_time_cd) IN ('3','03','4','04','5','05')
          AND etl_ymd BETWEEN '20260901' AND '20260930'
        GROUP BY 1,2,3,4,8
        """
    )
    con.execute(
        f"""
        CREATE OR REPLACE TABLE mode_dawn AS
        SELECT {common}, trim(move_trans) AS category_code
        FROM read_csv_auto('{_csv_sql_path(mode_glob)}', header=true,
                           all_varchar=true, union_by_name=true, filename=false)
        WHERE trim(st_time_cd) IN ('3','03','4','04','5','05')
          AND etl_ymd BETWEEN '20260901' AND '20260930'
        GROUP BY 1,2,3,4,8
        """
    )
    con.execute(
        """
        CREATE OR REPLACE VIEW purpose_core AS
        SELECT p.* FROM purpose_dawn p
        JOIN seoul_codes o ON p.origin_id=o.code
        JOIN seoul_codes d ON p.destination_id=d.code
        """
    )
    con.execute(
        """
        CREATE OR REPLACE VIEW mode_core AS
        SELECT p.* FROM mode_dawn p
        JOIN seoul_codes o ON p.origin_id=o.code
        JOIN seoul_codes d ON p.destination_id=d.code
        """
    )


def _mobility_outputs(con: duckdb.DuckDBPyConnection, out: Path) -> dict[str, int]:
    rows: dict[str, int] = {}
    day_counts = "CASE day_type WHEN 'weekday' THEN 22 ELSE 8 END"
    rows["od_hourly.csv"] = _export_query(
        con,
        out / "od_hourly.csv",
        f"""
        SELECT origin_id, destination_id, hour, day_type,
               round(sum(demand_value),2) AS estimated_movements,
               round(sum(demand_value)/({day_counts}),4) AS daily_average,
               sum(source_rows) AS source_rows,
               sum(nonmissing_cnt_rows) AS nonmissing_cnt_rows
        FROM purpose_core GROUP BY 1,2,3,4 ORDER BY 5 DESC,1,2,3,4
        """,
    )
    rows["hour_summary.csv"] = _export_query(
        con,
        out / "hour_summary.csv",
        f"""
        SELECT hour, day_type, round(sum(demand_value),2) AS estimated_movements,
               round(sum(demand_value)/({day_counts}),4) AS daily_average,
               sum(source_rows) AS source_rows,
               sum(nonmissing_cnt_rows) AS nonmissing_cnt_rows
        FROM purpose_core GROUP BY 1,2 ORDER BY 1,2
        """,
    )
    for side in ("origin", "destination"):
        ident = f"{side}_id"
        rows[f"{side}_concentration.csv"] = _export_query(
            con,
            out / f"{side}_concentration.csv",
            f"""
            WITH a AS (
              SELECT {ident}, hour, sum(demand_value) AS v
              FROM purpose_core GROUP BY 1,2
            )
            SELECT a.{ident}, c.full_name, c.district_name, a.hour,
                   round(a.v,2) AS estimated_movements,
                   dense_rank() OVER (PARTITION BY hour ORDER BY a.v DESC) AS concentration_rank,
                   round(a.v/sum(a.v) OVER (PARTITION BY hour),8) AS hour_share
            FROM a JOIN seoul_codes c ON a.{ident}=c.code
            ORDER BY hour, concentration_rank, {ident}
            """,
        )
    rows["destination_purpose.csv"] = _export_query(
        con,
        out / "destination_purpose.csv",
        """
        WITH a AS (
          SELECT destination_id, hour, category_code, sum(demand_value) AS v
          FROM purpose_core GROUP BY 1,2,3
        )
        SELECT a.destination_id, c.full_name, c.district_name, a.hour,
               a.category_code AS purpose_code,
               CASE a.category_code WHEN '1' THEN 'commuting' WHEN '2' THEN 'school'
                 WHEN '3' THEN 'home' WHEN '4' THEN 'shopping' WHEN '5' THEN 'tourism'
                 WHEN '6' THEN 'hospital' WHEN '7' THEN 'other' ELSE 'unknown' END AS purpose_label,
               round(a.v,2) AS estimated_movements,
               dense_rank() OVER (PARTITION BY a.destination_id,a.hour ORDER BY a.v DESC) AS purpose_rank,
               round(a.v/sum(a.v) OVER (PARTITION BY a.destination_id,a.hour),8) AS destination_hour_share
        FROM a JOIN seoul_codes c ON a.destination_id=c.code
        ORDER BY destination_id,hour,purpose_rank,purpose_code
        """,
    )
    for axis, view, label_expr in (
        ("purpose", "purpose_core", "CASE category_code WHEN '1' THEN 'commuting' WHEN '2' THEN 'school' WHEN '3' THEN 'home' WHEN '4' THEN 'shopping' WHEN '5' THEN 'tourism' WHEN '6' THEN 'hospital' WHEN '7' THEN 'other' ELSE 'unknown' END"),
        ("mode", "mode_core", "CASE category_code WHEN '1' THEN 'air' WHEN '2' THEN 'rail' WHEN '3' THEN 'express_bus' WHEN '4' THEN 'metropolitan_bus' WHEN '5' THEN 'ordinary_bus' WHEN '6' THEN 'subway' WHEN '7' THEN 'walking' WHEN '8' THEN 'vehicle' WHEN '9' THEN 'other' ELSE 'unknown' END"),
    ):
        rows[f"{axis}_composition.csv"] = _export_query(
            con,
            out / f"{axis}_composition.csv",
            f"""
            WITH a AS (SELECT hour, category_code, sum(demand_value) AS v FROM {view} GROUP BY 1,2)
            SELECT hour, category_code AS {axis}_code, {label_expr} AS {axis}_label,
                   round(v,2) AS estimated_movements,
                   round(v/sum(v) OVER (PARTITION BY hour),8) AS hour_share
            FROM a ORDER BY hour, estimated_movements DESC, category_code
            """,
        )
    rows["boundary_flow_summary.csv"] = _export_query(
        con,
        out / "boundary_flow_summary.csv",
        f"""
        WITH a AS (
          SELECT p.day_type,p.hour,
            CASE WHEN o.code IS NOT NULL AND d.code IS NOT NULL THEN 'internal'
                 WHEN o.code IS NULL AND d.code IS NOT NULL THEN 'inbound'
                 WHEN o.code IS NOT NULL AND d.code IS NULL THEN 'outbound'
                 ELSE 'outside_or_unmapped' END AS flow_scope,
            sum(p.demand_value) AS v, sum(p.source_rows) AS source_rows
          FROM purpose_dawn p LEFT JOIN seoul_codes o ON p.origin_id=o.code
          LEFT JOIN seoul_codes d ON p.destination_id=d.code GROUP BY 1,2,3
        )
        SELECT day_type,hour,flow_scope,round(v,2) AS estimated_movements,
               round(v/({day_counts}),4) AS daily_average,source_rows
        FROM a ORDER BY hour,day_type,flow_scope
        """,
    )
    rows["map_dong_demand.csv"] = _export_query(
        con,
        out / "map_dong_demand.csv",
        """
        WITH o AS (SELECT origin_id AS code,sum(demand_value) AS v FROM purpose_core GROUP BY 1),
             d AS (SELECT destination_id AS code,sum(demand_value) AS v FROM purpose_core GROUP BY 1)
        SELECT c.code AS admi_cd,c.full_name,c.district_name,
               round(o.v,2) AS origin_estimated_movements,
               round(d.v,2) AS destination_estimated_movements
        FROM seoul_codes c LEFT JOIN o USING(code) LEFT JOIN d USING(code)
        WHERE o.v IS NOT NULL OR d.v IS NOT NULL ORDER BY c.code
        """,
    )
    rows["demand_candidates.csv"] = _export_query(
        con,
        out / "demand_candidates.csv",
        """
        WITH a AS (
          SELECT origin_id,destination_id,hour,sum(demand_value) AS v,
                 sum(source_rows) AS source_rows,sum(nonmissing_cnt_rows) AS known_rows
          FROM purpose_core GROUP BY 1,2,3
        ), ranked AS (
          SELECT *,row_number() OVER (ORDER BY v DESC,origin_id,destination_id,hour) AS rnk
          FROM a WHERE v IS NOT NULL AND origin_id <> destination_id
        )
        SELECT printf('DMD-%04d',rnk) AS candidate_id,origin_id,destination_id,
               printf('%02d:00',hour) AS time_bin_start,printf('%02d:00',hour+1) AS time_bin_end,
               round(v,2) AS demand_value,'estimated movements (monthly sum)' AS demand_unit,
               rnk AS demand_rank,'top 100 distinct-dong internal Seoul OD-hour by OA-22300 weekday plus weekend monthly estimated movements' AS selection_rule,
               CASE WHEN known_rows < source_rows THEN 'cnt_missing_rows_excluded_from_sum; suppressed/unpublished cells unknown'
                    ELSE 'suppressed/unpublished cells unknown; aggregate estimate, not unique persons' END AS uncertainty_flag,
               'OA-22300' AS source_dataset_ids,
               o.full_name AS origin_name,o.district_name AS origin_district_name,
               d.full_name AS destination_name,d.district_name AS destination_district_name
        FROM ranked JOIN seoul_codes o ON ranked.origin_id=o.code
        JOIN seoul_codes d ON ranked.destination_id=d.code
        WHERE rnk <= 100 ORDER BY rnk
        """,
    )
    rows["intrazonal_summary.csv"] = _export_query(
        con,
        out / "intrazonal_summary.csv",
        f"""
        WITH a AS (
          SELECT hour,day_type,
                 sum(CASE WHEN origin_id=destination_id THEN demand_value END) AS intrazonal,
                 sum(demand_value) AS total
          FROM purpose_core GROUP BY 1,2
        )
        SELECT hour,day_type,round(intrazonal,2) AS intrazonal_estimated_movements,
               round(intrazonal/({day_counts}),4) AS intrazonal_daily_average,
               round(total,2) AS all_internal_estimated_movements,
               round(intrazonal/total,8) AS intrazonal_share
        FROM a ORDER BY hour,day_type
        """,
    )
    return rows


def _parse_number(value: str) -> float | None:
    text = value.strip().replace(",", "")
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _sum_known(values: Iterable[float | None]) -> float | None:
    vals = [v for v in values if v is not None]
    return sum(vals) if vals else None


def _bus_outputs(root: Path, out: Path) -> dict[str, int]:
    source = next((root / "data/raw/bus/OA-12913").glob("*.csv"))
    records: list[dict[str, object]] = []
    missing_cells = 0
    with source.open("r", encoding="cp949", newline="") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError("OA-12913 has no header")
        names = reader.fieldnames
        # The audited source columns are stable; identify hour columns by their leading hour.
        board_cols = [n for n in names if any(n.startswith(f"{h}시") for h in (3, 4, 5)) and "승차" in n]
        alight_cols = [n for n in names if any(n.startswith(f"{h}시") for h in (3, 4, 5)) and "하차" in n]
        if len(board_cols) != 3 or len(alight_cols) != 3:
            raise ValueError(f"Unexpected OA-12913 dawn columns: {board_cols!r}, {alight_cols!r}")
        fixed = {
            "route_no": "노선번호",
            "route_name": "노선명",
            "stop_id": "표준버스정류장ID",
            "ars_id": "버스정류장ARS번호",
            "stop_name": "역명",
            "transport_type_code": "교통수단타입코드",
            "transport_type_name": "교통수단타입명",
        }
        for row in reader:
            for hour, bcol, acol in zip((3, 4, 5), board_cols, alight_cols):
                board = _parse_number(row[bcol])
                alight = _parse_number(row[acol])
                missing_cells += int(board is None) + int(alight is None)
                records.append(
                    {**{key: row.get(col, "").strip() for key, col in fixed.items()},
                     "hour": hour, "boardings": board, "alightings": alight}
                )

    def aggregate(keys: tuple[str, ...]) -> list[dict[str, object]]:
        grouped: dict[tuple[object, ...], dict[str, object]] = {}
        for rec in records:
            key = tuple(rec[k] for k in keys)
            agg = grouped.setdefault(key, {k: rec[k] for k in keys} | {"b": [], "a": [], "rows": 0})
            agg["b"].append(rec["boardings"])
            agg["a"].append(rec["alightings"])
            agg["rows"] += 1
        result = []
        for agg in grouped.values():
            result.append({k: agg[k] for k in keys} | {
                "boardings": _sum_known(agg["b"]), "alightings": _sum_known(agg["a"]),
                "source_rows": agg["rows"]})
        return result

    rows: dict[str, int] = {}
    hour = aggregate(("hour",))
    rows["bus_hour_summary.csv"] = _write_csv(
        out / "bus_hour_summary.csv", ["hour", "boardings", "alightings", "source_rows"],
        ([r["hour"], r["boardings"], r["alightings"], r["source_rows"]] for r in sorted(hour, key=lambda x: x["hour"])),
    )
    transport = aggregate(("transport_type_code", "transport_type_name", "hour"))
    transport.sort(key=lambda r: (r["hour"], -(r["boardings"] or -1), r["transport_type_code"]))
    rows["bus_transport_type_hour.csv"] = _write_csv(
        out / "bus_transport_type_hour.csv",
        ["transport_type_code", "transport_type_name", "hour", "boardings", "alightings", "source_rows"],
        ([r[k] for k in ("transport_type_code", "transport_type_name", "hour", "boardings", "alightings", "source_rows")] for r in transport),
    )
    routes = aggregate(("route_no", "route_name", "hour"))
    routes.sort(key=lambda r: (r["hour"], -(r["boardings"] or -1), r["route_no"]))
    rows["bus_route_concentration.csv"] = _write_csv(
        out / "bus_route_concentration.csv",
        ["route_no", "route_name", "hour", "boardings", "alightings", "source_rows"],
        ([r[k] for k in ("route_no", "route_name", "hour", "boardings", "alightings", "source_rows")] for r in routes),
    )
    stop_names: dict[tuple[str, str], str] = {}
    for rec in records:
        key = (str(rec["stop_id"]), str(rec["ars_id"]))
        name = str(rec["stop_name"])
        if key not in stop_names or name < stop_names[key]:
            stop_names[key] = name
    stops = aggregate(("stop_id", "ars_id", "hour"))
    stops.sort(key=lambda r: (r["hour"], -((r["boardings"] or 0) + (r["alightings"] or 0)), r["stop_id"]))
    ranked: list[list[object]] = []
    rank_by_hour = defaultdict(int)
    for rec in stops:
        rank_by_hour[rec["hour"]] += 1
        rank = rank_by_hour[rec["hour"]]
        if rank <= 100:
            ranked.append([rec["stop_id"], rec["ars_id"], stop_names[(str(rec["stop_id"]), str(rec["ars_id"]))], rec["hour"], rec["boardings"], rec["alightings"], rank])
    rows["bus_stop_concentration.csv"] = _write_csv(
        out / "bus_stop_concentration.csv",
        ["stop_id", "ars_id", "stop_name", "hour", "boardings", "alightings", "activity_rank_within_hour"], ranked,
    )
    # OA-12913 has no coordinates. Attach the independently audited OA-1095 snapshot
    # by standard stop ID; ambiguous duplicate coordinates remain explicit.
    from openpyxl import load_workbook

    oa1095 = next((root / "data/raw/bus/OA-1095").glob("*.xlsx"))
    workbook = load_workbook(oa1095, read_only=True, data_only=True)
    sheet = workbook.active
    source_rows = sheet.iter_rows(values_only=True)
    header = [str(value) if value is not None else "" for value in next(source_rows)]
    col = {name: i for i, name in enumerate(header)}
    coords: dict[str, set[tuple[object, object]]] = defaultdict(set)
    for source_row in source_rows:
        stop_id = str(source_row[col["NODE_ID"]]).strip() if source_row[col["NODE_ID"]] is not None else ""
        if stop_id:
            coords[stop_id].add((source_row[col["X좌표"]], source_row[col["Y좌표"]]))
    workbook.close()

    all_stops = aggregate(("stop_id", "ars_id"))
    all_stops.sort(key=lambda r: (-((r["boardings"] or 0) + (r["alightings"] or 0)), r["stop_id"]))
    map_rows = []
    for rec in all_stops:
        matches = sorted(coords.get(str(rec["stop_id"]), set()), key=lambda pair: (str(pair[0]), str(pair[1])))
        lon = matches[0][0] if len(matches) == 1 else None
        lat = matches[0][1] if len(matches) == 1 else None
        stop_name = stop_names[(str(rec["stop_id"]), str(rec["ars_id"]))]
        map_rows.append([rec["stop_id"], rec["ars_id"], stop_name, lon, lat,
                         "matched_unique" if len(matches) == 1 else ("ambiguous" if matches else "unmatched"),
                         rec["boardings"], rec["alightings"]])
    rows["bus_stop_map.csv"] = _write_csv(
        out / "bus_stop_map.csv",
        ["stop_id", "ars_id", "stop_name", "longitude", "latitude", "coordinate_status", "boardings", "alightings"],
        map_rows,
    )
    candidates = []
    for rank, rec in enumerate(all_stops[:100], 1):
        matches = sorted(coords.get(str(rec["stop_id"]), set()), key=lambda pair: (str(pair[0]), str(pair[1])))
        lon = matches[0][0] if len(matches) == 1 else None
        lat = matches[0][1] if len(matches) == 1 else None
        candidates.append([
            f"BUS-{rank:04d}", rec["stop_id"], rec["ars_id"], stop_names[(str(rec["stop_id"]), str(rec["ars_id"]))], lon, lat,
            rec["boardings"], rec["alightings"], rank,
            "top 100 stops by OA-12913 03-05 monthly boarding plus alighting activity",
            "monthly totals only; no weekday/weekend axis; activity is not unmet demand", "OA-12913|OA-1095",
        ])
    rows["bus_demand_candidates.csv"] = _write_csv(
        out / "bus_demand_candidates.csv",
        ["candidate_id", "stop_id", "ars_id", "stop_name", "longitude", "latitude", "boardings", "alightings", "activity_rank", "selection_rule", "uncertainty_flag", "source_dataset_ids"],
        candidates,
    )
    rows["_bus_missing_cells"] = missing_cells
    rows["_bus_boardings"] = int(sum((r["boardings"] or 0) for r in hour))
    rows["_bus_alightings"] = int(sum((r["alightings"] or 0) for r in hour))
    rows["_bus_mapped_stops"] = sum(row[5] == "matched_unique" for row in map_rows)
    rows["_bus_unmatched_stops"] = sum(row[5] == "unmatched" for row in map_rows)
    rows["_bus_ambiguous_stops"] = sum(row[5] == "ambiguous" for row in map_rows)
    return rows


def run(root: Path, memory_limit: str = "4GB", threads: int = 2) -> dict[str, object]:
    root = root.resolve()
    out = root / "outputs/intermediate/demand"
    out.mkdir(parents=True, exist_ok=True)
    scratch = out / ".scratch.duckdb"
    temp_dir = out / ".duckdb_tmp"
    temp_dir.mkdir(exist_ok=True)
    con = duckdb.connect(str(scratch))
    con.execute(f"SET memory_limit='{memory_limit}'")
    con.execute(f"SET threads={int(threads)}")
    con.execute(f"SET temp_directory='{_csv_sql_path(temp_dir)}'")
    con.execute("SET preserve_insertion_order=false")
    _build_mobility_aggregates(con, root)
    output_rows = _mobility_outputs(con, out)
    purpose_total = con.execute("SELECT round(sum(demand_value),2) FROM purpose_core").fetchone()[0]
    mode_total = con.execute("SELECT round(sum(demand_value),2) FROM mode_core").fetchone()[0]
    grouped_counts = {
        "purpose_dawn": con.execute("SELECT count(*) FROM purpose_dawn").fetchone()[0],
        "purpose_core": con.execute("SELECT count(*) FROM purpose_core").fetchone()[0],
        "mode_dawn": con.execute("SELECT count(*) FROM mode_dawn").fetchone()[0],
        "mode_core": con.execute("SELECT count(*) FROM mode_core").fetchone()[0],
    }
    con.close()
    bus_rows = _bus_outputs(root, out)
    bus_missing = bus_rows.pop("_bus_missing_cells")
    bus_totals = {
        "boardings_03_05_monthly": bus_rows.pop("_bus_boardings"),
        "alightings_03_05_monthly": bus_rows.pop("_bus_alightings"),
        "coordinate_matched_unique_stops": bus_rows.pop("_bus_mapped_stops"),
        "coordinate_unmatched_stops": bus_rows.pop("_bus_unmatched_stops"),
        "coordinate_ambiguous_stops": bus_rows.pop("_bus_ambiguous_stops"),
    }
    output_rows.update(bus_rows)
    for candidate_file in (out / "demand_candidates.csv", out / "bus_demand_candidates.csv"):
        with candidate_file.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if candidate_file.name == "demand_candidates.csv" and (
                reader.fieldnames is None or not set(REQUIRED_CANDIDATE_COLUMNS).issubset(reader.fieldnames)
            ):
                raise ValueError(f"Candidate schema mismatch: {reader.fieldnames}")
    artifacts = {}
    for path in sorted(out.glob("*.csv")):
        artifacts[path.name] = {"rows": output_rows[path.name], "sha256": _sha256(path)}
    summary = {
        "status": "complete_limited",
        "readiness_mode": "separate_only",
        "analysis_window": "departure 03:00 inclusive to 06:00 exclusive, Asia/Seoul",
        "core_geography": "both origin and destination exactly matched to ADMI_202608 Seoul codes",
        "day_type_days": {"weekday": 22, "weekend": 8},
        "purpose_axis_internal_total": purpose_total,
        "mode_axis_internal_total": mode_total,
        "axis_rule": "Purpose and mode are alternative classifications and are never added or subtracted.",
        "bus_rule": "OA-12913 monthly boarding/alighting is independently described and never subtracted from mobility counts.",
        "bus_missing_dawn_cells": bus_missing,
        "bus_totals": bus_totals,
        "grouped_rows": grouped_counts,
        "artifacts": artifacts,
        "unsupported": ["bus weekday/weekend comparison", "service gap", "unmet demand", "unique persons"],
    }
    (out / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lineage = {
        "sources": {
            "OA-22300": "30 daily purpose CSVs, September 2026",
            "OA-22657": "30 daily mode CSVs, September 2026",
            "OA-12913": "September 2026 route-stop-hour monthly totals",
            "OA-1095": "September 2, 2026 route-stop snapshot used only for stop coordinates",
            "mobility_codebook": "ADMI_202608.csv",
        },
        "filters": ["etl_ymd 20260901..20260930", "st_time_cd exactly one of 3,03,4,04,5,05", "core OD has both codes in Seoul ADMI_202608"],
        "calculations": {
            "monthly_sum": "sum(cnt) over retained aggregate rows; missing cnt is excluded, never replaced by zero",
            "daily_average": "weekday monthly subtotal / 22; weekend monthly subtotal / 8",
            "concentration_share": "dong estimated movement sum / all internal-Seoul dong sum within hour",
            "candidate_rule": "top 100 distinct-dong internal-Seoul OD-hour cells by OA-22300 weekday plus weekend monthly estimated movement sum; intrazonal cells retained only in intrazonal_summary.csv",
            "bus_candidate_rule": "top 100 stops by sum of OA-12913 03,04,05 monthly boardings plus alightings",
        },
        "non_additivity": "OA-22300 purpose and OA-22657 mode totals are parallel axes; OA-12913 has a different unit and is not differenced.",
    }
    (out / "lineage.json").write_text(json.dumps(lineage, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    notes = """# Demand analysis method notes

The core result uses estimated movement records whose origin and destination both match the August 2026 Seoul administrative-dong codebook. Departure codes are matched only to `3|03|4|04|5|05`; HH and HHMM codes are not conflated. Calendar Monday-Friday and Saturday-Sunday contain 22 and 8 observed days respectively, without holiday adjustment.

`cnt` is summed as estimated movements. It is neither a unique-person count nor observed unmet demand. Missing cells are not replaced by zero. Purpose and mode files are alternative classifications and are reported independently. OA-12913 boarding/alighting monthly totals use a different measurement unit and are described independently; no subtraction from mobility estimates is performed. Its lack of a daily axis makes weekday/weekend comparison unsupported.

Demand candidates are the 100 highest distinct-dong internal-Seoul OD-hour cells by OA-22300 weekday-plus-weekend monthly estimated movements. Same-dong movements are retained in `intrazonal_summary.csv`, but are not used as corridor candidates. Bus candidates are the 100 stops with the largest combined 03:00-05:59 monthly boarding and alighting activity. OA-1095 supplies map coordinates where a standard stop ID has one unique coordinate pair. These are demand-side screening candidates only and contain no first-service, transfer, walking, service-calendar, service-gap, or causal conclusion.

No predictive or clustering model is used because the requested descriptive rankings and compositions are directly identifiable from the available aggregates, while a model would add assumptions without a validated outcome target.
"""
    (out / "method_notes.md").write_text(notes, encoding="utf-8")
    # Scratch data is a reproducibility implementation detail, not an output artifact.
    if scratch.exists():
        scratch.unlink()
    if temp_dir.exists():
        shutil.rmtree(temp_dir)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--memory-limit", default="4GB")
    parser.add_argument("--threads", type=int, default=2)
    args = parser.parse_args()
    result = run(args.root, args.memory_limit, args.threads)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

